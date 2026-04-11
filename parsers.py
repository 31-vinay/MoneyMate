import csv
import io
import re
import ssl
import imaplib
import email as email_lib
from email.header import decode_header
from email.utils import parsedate_to_datetime
from datetime import datetime, timedelta

import openpyxl
import xlrd
import msoffcrypto
import pdfplumber
from bs4 import BeautifulSoup


_DATE_FMTS = [
    "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y",
    "%Y-%m-%d", "%m/%d/%Y", "%d %b %Y", "%d %b %y",
    "%d-%b-%Y", "%d-%b-%y", "%d/%b/%Y", "%d/%b/%y",
]


def _parse_date(s):
    s = s.strip()
    for fmt in _DATE_FMTS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _clean_amount(s):
    s = re.sub(r"[₹,\s]", "", str(s)).strip()
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _make_txn(date_obj, desc, amount, txn_type):
    return {
        "date": date_obj.strftime("%Y-%m-%d"),
        "description": desc.strip()[:200],
        "amount": round(amount, 2),
        "type": txn_type,
        "sub_cat": "Bank Import",
        "main_cat": "Other Expenses" if txn_type == "expense" else "Other Income",
        "subject": desc.strip()[:200],
    }


def _detect_columns(headers):
    h = [str(x).lower().strip() for x in headers]
    date_kw = ["date", "txn date", "value date", "transaction date", "posting date"]
    desc_kw = ["description", "narration", "particulars", "remarks", "details",
               "transaction remarks", "transaction details"]
    debit_kw = ["debit", "withdrawal", "dr", "debit amount", "withdrawal amount", "amount (dr)"]
    credit_kw = ["credit", "deposit", "cr", "credit amount", "deposit amount", "amount (cr)"]
    amount_kw = ["amount", "net amount"]

    def find(kws):
        for k in kws:
            for i, hh in enumerate(h):
                if k in hh:
                    return i
        return None

    di = find(date_kw)
    ni = find(desc_kw)
    dbi = find(debit_kw)
    cri = find(credit_kw)
    ami = find(amount_kw) if dbi is None else None
    if di is None or ni is None:
        return None
    return di, ni, dbi, cri, ami


def _rows_to_txns(rows, header_idx):
    headers = rows[header_idx]
    result = _detect_columns(headers)
    if result is None:
        return []
    di, ni, dbi, cri, ami = result
    txns = []
    for row in rows[header_idx + 1:]:
        max_col = max(x for x in [di, ni, dbi, cri, ami] if x is not None)
        if len(row) <= max_col:
            continue
        date_val = _parse_date(str(row[di]))
        if date_val is None:
            continue
        desc = str(row[ni]).strip()
        if not desc or desc.lower() in ("", "nan", "none"):
            continue
        debit = _clean_amount(row[dbi]) if dbi is not None and dbi < len(row) else None
        credit = _clean_amount(row[cri]) if cri is not None and cri < len(row) else None
        amount = _clean_amount(row[ami]) if ami is not None and ami < len(row) else None
        if debit and debit > 0:
            txns.append(_make_txn(date_val, desc, debit, "expense"))
        if credit and credit > 0:
            txns.append(_make_txn(date_val, desc, credit, "income"))
        if amount and dbi is None and cri is None:
            t = "income" if amount > 0 else "expense"
            txns.append(_make_txn(date_val, desc, abs(amount), t))
    return txns


def _decrypt_office_file(file_bytes, password):
    enc_file = io.BytesIO(file_bytes)
    dec_file = io.BytesIO()
    office_file = msoffcrypto.OfficeFile(enc_file)
    office_file.load_key(password=password)
    office_file.decrypt(dec_file)
    dec_file.seek(0)
    return dec_file.read()


def _is_encrypted_office(file_bytes):
    try:
        enc_file = io.BytesIO(file_bytes)
        office_file = msoffcrypto.OfficeFile(enc_file)
        return office_file.is_encrypted()
    except Exception:
        return False


def parse_bank_statement_csv(content_bytes):
    text = content_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [r for r in reader if any(c.strip() for c in r)]
    for i, row in enumerate(rows):
        if _detect_columns(row):
            return _rows_to_txns(rows, i)
    return []


def parse_bank_statement_xlsx(content_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c) if c is not None else "" for c in row]
        if any(c.strip() for c in cells):
            rows.append(cells)
    for i, row in enumerate(rows):
        if _detect_columns(row):
            return _rows_to_txns(rows, i)
    return []


def parse_bank_statement_xls(content_bytes):
    wb = xlrd.open_workbook(file_contents=content_bytes)
    ws = wb.sheet_by_index(0)
    rows = []
    for rx in range(ws.nrows):
        cells = []
        for cx in range(ws.ncols):
            cell = ws.cell(rx, cx)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    dt = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                    cells.append(dt.strftime("%d/%m/%Y"))
                except Exception:
                    cells.append(str(cell.value))
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                v = cell.value
                cells.append(str(int(v)) if v == int(v) else str(v))
            else:
                cells.append(str(cell.value).strip())
        if any(c.strip() for c in cells):
            rows.append(cells)
    for i, row in enumerate(rows):
        if _detect_columns(row):
            return _rows_to_txns(rows, i)
    return []


def parse_bank_statement_pdf(content_bytes):
    DATE_RE = re.compile(
        r"\b(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4}|\d{1,2}[/\-][A-Za-z]{3}[/\-]\d{2,4})\b"
    )
    AMOUNT_RE = re.compile(r"(\d{1,3}(?:,\d{2,3})*(?:\.\d{2})?)")
    txns = []
    with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                if not table:
                    continue
                rows = [[str(c) if c else "" for c in row] for row in table]
                for i, row in enumerate(rows):
                    if _detect_columns(row):
                        txns.extend(_rows_to_txns(rows, i))
                        break
            if txns:
                continue
            text = page.extract_text() or ""
            for line in text.splitlines():
                line = line.strip()
                dates = DATE_RE.findall(line)
                if not dates:
                    continue
                date_obj = _parse_date(dates[0])
                if not date_obj:
                    continue
                amounts = AMOUNT_RE.findall(line)
                if not amounts:
                    continue
                desc = DATE_RE.sub("", line)
                for a in amounts:
                    desc = desc.replace(a, "")
                desc = re.sub(r"\s+", " ", desc).strip(" ,.-/")
                if not desc:
                    desc = "Bank transaction"
                amt = _clean_amount(amounts[-2]) if len(amounts) >= 2 else _clean_amount(amounts[-1])
                if amt and amt > 0:
                    low = line.lower()
                    if re.search(r"\bcr\b|\bcredit\b|\bdeposit\b", low):
                        txns.append(_make_txn(date_obj, desc, amt, "income"))
                    else:
                        txns.append(_make_txn(date_obj, desc, amt, "expense"))
    seen = set()
    unique = []
    for t in txns:
        key = (t["date"], t["amount"], t["type"])
        if key not in seen:
            seen.add(key)
            unique.append(t)
    return unique


def parse_bank_statement(file_bytes, filename, password=None):
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return parse_bank_statement_pdf(file_bytes)
    elif ext in ("xlsx", "xls"):
        data = file_bytes
        if password:
            try:
                data = _decrypt_office_file(file_bytes, password)
            except Exception:
                raise ValueError("Wrong password or unable to decrypt the file.")
        elif _is_encrypted_office(file_bytes):
            raise ValueError(
                "This file is password-protected. Please enter the statement password."
            )
        if ext == "xls":
            return parse_bank_statement_xls(data)
        else:
            return parse_bank_statement_xlsx(data)
    elif ext == "csv":
        return parse_bank_statement_csv(file_bytes)
    return []


IMAP_PRESETS = {
    "gmail":   {"host": "imap.gmail.com",        "port": 993},
    "outlook": {"host": "imap-mail.outlook.com", "port": 993},
    "yahoo":   {"host": "imap.mail.yahoo.com",   "port": 993},
    "hotmail": {"host": "imap-mail.outlook.com", "port": 993},
    "icloud":  {"host": "imap.mail.me.com",      "port": 993},
}

FINANCIAL_SUBJECT_KEYWORDS = [
    "transaction", "payment", "purchase", "receipt", "order", "invoice",
    "debit", "credit", "charged", "statement", "bill", "transfer",
    "alert", "notification", "confirmation", "refund", "deposit",
    "subscription", "auto-pay", "autopay", "due", "amount",
]

FINANCIAL_SENDER_KEYWORDS = [
    "bank", "paypal", "paytm", "stripe", "amazon", "netflix", "spotify",
    "apple", "google", "microsoft", "hulu", "prime", "uber", "lyft",
    "razorpay", "hdfc", "sbi", "icici", "axis", "netsuite", "venmo",
    "cashapp", "zelle", "chase", "citibank", "wells", "fargo",
]

CATEGORY_KEYWORD_MAP = [
    (["grocery", "supermarket", "safeway", "kroger", "walmart", "costco", "whole foods", "amazon fresh", "trader joe"], ("Food & Groceries", "Groceries")),
    (["restaurant", "dining", "bistro", "cafe", "diner", "eatery", "sushi", "pizza", "burger"], ("Food & Groceries", "Dining Out")),
    (["coffee", "starbucks", "dunkin", "costa"], ("Food & Groceries", "Coffee Shops")),
    (["food delivery", "doordash", "grubhub", "ubereats", "zomato", "swiggy"], ("Food & Groceries", "Food Delivery")),
    (["fast food", "mcdonald", "kfc", "subway", "domino", "taco bell", "wendy", "burger king"], ("Food & Groceries", "Fast Food")),
    (["uber", "lyft", "taxi", "rideshare", "ola", "grab"], ("Transportation", "Taxi/Rideshare")),
    (["fuel", "gas station", "petrol", "shell", "bp ", "chevron", "exxon"], ("Transportation", "Fuel")),
    (["metro", "bus", "transit", "train", "subway pass", "rail"], ("Transportation", "Public Transport")),
    (["netflix", "hulu", "disney+", "hbo", "prime video", "apple tv", "peacock", "paramount"], ("Entertainment", "Streaming Services")),
    (["spotify", "apple music", "tidal", "deezer", "pandora", "youtube music"], ("Entertainment", "Music Streaming")),
    (["gym", "fitness", "planet fitness", "equinox", "crunch"], ("Personal & Lifestyle", "Gym Membership")),
    (["electricity", "electric", "power bill"], ("Utilities", "Electricity")),
    (["water bill"], ("Utilities", "Water")),
    (["internet", "broadband", "comcast", "xfinity", "att", "verizon", "spectrum"], ("Utilities", "Internet")),
    (["mobile", "phone bill", "t-mobile", "sprint", "cricket"], ("Utilities", "Mobile Phone")),
    (["amazon", "ebay", "etsy", "shopify", "online shopping", "shop", "purchase from"], ("Shopping", "Online Shopping")),
    (["doctor", "clinic", "hospital", "medical", "health", "dental", "pharmacy", "prescription"], ("Healthcare", "Doctor Visits")),
    (["insurance", "policy", "premium"], ("Insurance", "Health Insurance")),
    (["school", "tuition", "university", "college", "course", "udemy", "coursera"], ("Education", "School Tuition")),
    (["rent", "lease", "landlord"], ("Housing", "Rent")),
    (["salary", "payroll", "wages", "paycheck"], ("Income", "Salary")),
    (["refund", "cashback", "reward"], ("Income", "Refund")),
]


def _ei_decode_header(raw):
    parts = decode_header(raw or "")
    result = ""
    for part, enc in parts:
        if isinstance(part, bytes):
            result += part.decode(enc or "utf-8", errors="ignore")
        else:
            result += str(part)
    return result


def _ei_extract_text(msg):
    plain, html = "", ""
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            payload = part.get_payload(decode=True)
            if not payload:
                continue
            decoded = payload.decode("utf-8", errors="ignore")
            if ctype == "text/plain":
                plain += decoded
            elif ctype == "text/html":
                html += decoded
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            decoded = payload.decode("utf-8", errors="ignore")
            if msg.get_content_type() == "text/html":
                html = decoded
            else:
                plain = decoded
    if html and not plain.strip():
        soup = BeautifulSoup(html, "html.parser")
        plain = soup.get_text(separator=" ", strip=True)
    return plain


def _ei_parse_amount(text):
    patterns = [
        r'\$\s*([\d,]+\.?\d*)',
        r'USD\s+([\d,]+\.?\d*)',
        r'Rs\.?\s*([\d,]+\.?\d*)',
        r'INR\s+([\d,]+\.?\d*)',
        r'(?:amount|total|charged|debit|credit)[:\s]+(?:of\s+)?\$?\s*([\d,]+\.?\d*)',
        r'payment of\s+\$?\s*([\d,]+\.?\d*)',
        r'\b([\d,]{1,10}\.\d{2})\b',
    ]
    for pat in patterns:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            try:
                val = float(m.group(1).replace(",", ""))
                if 0.01 <= val <= 999999:
                    return round(val, 2)
            except ValueError:
                continue
    return None


def _ei_guess_category(subject, body):
    combined = (subject + " " + body).lower()
    for keywords, (main_cat, sub_cat) in CATEGORY_KEYWORD_MAP:
        if any(kw in combined for kw in keywords):
            return main_cat, sub_cat
    return "Other Expenses", "Miscellaneous"


def _ei_is_income(subject, body):
    combined = (subject + " " + body).lower()
    income_signals = [
        "received", "credited to your account", "deposit", "refund", "cashback",
        "salary", "payroll", "transfer received", "payment received", "reward",
    ]
    return any(s in combined for s in income_signals)


def _ei_is_financial(subject, from_addr):
    sub_lower = subject.lower()
    from_lower = from_addr.lower()
    if any(kw in sub_lower for kw in FINANCIAL_SUBJECT_KEYWORDS):
        return True
    if any(kw in from_lower for kw in FINANCIAL_SENDER_KEYWORDS):
        return True
    return False


def scan_imap_emails(host, port, email_addr, password, days=30):
    ctx = ssl.create_default_context()
    mail = imaplib.IMAP4_SSL(host, int(port), ssl_context=ctx)
    mail.login(email_addr, password)
    mail.select("INBOX")

    since_date = (datetime.utcnow() - timedelta(days=days)).strftime("%d-%b-%Y")
    _, message_ids = mail.search(None, f"SINCE {since_date}")
    msg_ids = message_ids[0].split()
    msg_ids = msg_ids[-300:][::-1]

    transactions = []
    seen_ids = set()

    for mid in msg_ids:
        try:
            _, msg_data = mail.fetch(mid, "(RFC822)")
            raw = msg_data[0][1]
            msg = email_lib.message_from_bytes(raw)

            subject = _ei_decode_header(msg.get("Subject", ""))
            from_addr = msg.get("From", "")
            date_str = msg.get("Date", "")
            msg_id = msg.get("Message-ID", mid.decode()).strip("<> ")

            if msg_id in seen_ids:
                continue
            seen_ids.add(msg_id)

            if not _ei_is_financial(subject, from_addr):
                continue

            body = _ei_extract_text(msg)
            amount = _ei_parse_amount(subject + " " + body[:3000])
            if not amount:
                continue

            is_income = _ei_is_income(subject, body)
            main_cat, sub_cat = _ei_guess_category(subject, body)
            if is_income:
                main_cat, sub_cat = "Income", "Transfer/Other"

            try:
                txn_date = parsedate_to_datetime(date_str).strftime("%Y-%m-%d")
            except Exception:
                txn_date = datetime.utcnow().strftime("%Y-%m-%d")

            transactions.append({
                "msg_id": msg_id,
                "subject": subject[:120],
                "from": from_addr[:100],
                "date": txn_date,
                "amount": amount,
                "type": "income" if is_income else "expense",
                "main_cat": main_cat,
                "sub_cat": sub_cat,
                "description": subject[:200],
            })
        except Exception:
            continue

    mail.logout()
    return transactions
